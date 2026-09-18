"""U-18 — resource export to Word/PowerPoint (`app/services/export.py`).

Formats already-stored, already-judged content. No generation, no judge, no Mongo:
every resource dict here sets doc_id=None so _filename_stem/_source_line never call
pdf_store.get_document.
"""

from __future__ import annotations

from io import BytesIO
from unittest.mock import patch

import pytest

from app.services import export as export_service


def _resource(task: str, content, summary_style: str = None) -> dict:
    params = {"summary_style": summary_style} if summary_style else {}
    return {"doc_id": None, "task": task, "content": content, "params": params}


def test_docx_summary_narrative_keeps_paragraphs():
    from docx import Document

    resource = _resource("summary", "First point.\n\nSecond point.", summary_style="narrative")
    data = export_service._docx_bytes(resource)
    assert data[:2] == b"PK"
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "First point." in text
    assert "Second point." in text


def test_docx_summary_structured_uses_bullets():
    from docx import Document

    resource = _resource("summary", "- Point one\n- Point two", summary_style="structured")
    data = export_service._docx_bytes(resource)
    doc = Document(BytesIO(data))
    bullet_styles = [p.style.name for p in doc.paragraphs if "List" in p.style.name]
    assert bullet_styles, "expected at least one List Bullet paragraph"


def test_docx_mcq_lists_options_and_correct_answer():
    from docx import Document

    items = [{
        "question": "What is consideration?",
        "options": ["Price of the promise", "A gift", "A tort", "A crime"],
        "correct_answer": "Price of the promise",
    }]
    data = export_service._docx_bytes(_resource("mcq", items))
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "What is consideration?" in text
    assert "Price of the promise" in text


def test_pptx_mcq_builds_one_slide_per_question_plus_answer_key():
    from pptx import Presentation

    items = [
        {"question": "Q1?", "options": ["A", "B", "C", "D"], "correct_answer": "A"},
        {"question": "Q2?", "options": ["A", "B", "C", "D"], "correct_answer": "B"},
    ]
    data = export_service._pptx_bytes(_resource("mcq", items))
    assert data[:2] == b"PK"
    prs = Presentation(BytesIO(data))
    # cover slide + 2 question slides + answer key slide
    assert len(list(prs.slides)) == 4


def test_pptx_keypoints_groups_four_per_slide():
    from pptx import Presentation

    points = [f"Point {i}" for i in range(1, 10)]  # 9 points -> 3 slides of 4/4/1
    data = export_service._pptx_bytes(_resource("keypoints", points))
    prs = Presentation(BytesIO(data))
    # cover slide + 3 content slides
    assert len(list(prs.slides)) == 4


def test_export_resource_rejects_unknown_format():
    with patch.object(export_service.access, "require_resource",
                       return_value={"doc_id": None, "task": "summary", "content": "x", "params": {}}):
        with pytest.raises(ValueError, match="docx or pptx"):
            export_service.export_resource("u1", "r1", "xlsx")


def test_export_resource_defaults_to_docx_and_names_the_file():
    fake = {"doc_id": None, "task": "summary", "content": "Body text.", "params": {}}
    with patch.object(export_service.access, "require_resource", return_value=fake):
        data, media_type, filename = export_service.export_resource("u1", "r1", "")
    assert filename.endswith(".docx")
    assert media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert data[:2] == b"PK"
